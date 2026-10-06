#!/usr/bin/env python3
"""QA for the Northwind workshop pack. Asserts every planted fact, every
tie-out, text extraction of every invoice PDF, the forbidden-term scan, the
xlsx XML scan and the size budget. Exits non-zero on any failure.

Run from the repo root:
    uv run --with openpyxl --with fpdf2 --with pillow python3 scripts/qa_check.py
"""
import csv
import os
import re
import subprocess
import sys
import zipfile
from collections import defaultdict
from datetime import date

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate import INBOX_INVOICES, DUPLICATE_FILE, BLANK_SCAN, VENDORS  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D01, D02, D03, D04 = [os.path.join(ROOT, d) for d in ("01-invoices-inbox", "02-sources", "03-dashboard", "04-prospect")]
FAILS = []
PASSES = 0


def check(cond, msg):
    global PASSES
    if cond:
        PASSES += 1
        print("  ok   " + msg)
    else:
        FAILS.append(msg)
        print("  FAIL " + msg)


def pdftotext(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def read_csv(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def cents(x):
    return int(round(float(x) * 100))


print("QA: Northwind workshop pack")

# ------------------------------------------------------------ 02 sources
print("\n[02-sources]")
jobs = read_csv(os.path.join(D02, "jobs-export.csv"))
inv = read_csv(os.path.join(D02, "quickbooks-invoices-q3.csv"))
bills = read_csv(os.path.join(D02, "quickbooks-bills-paid-q3.csv"))
check(350 <= len(jobs) <= 520, "jobs-export has a modest row count (%d)" % len(jobs))
check(100 <= len(inv) <= 150, "quickbooks-invoices-q3 has ~120 rows (%d)" % len(inv))
check(40 <= len(bills) <= 50, "quickbooks-bills-paid-q3 has ~45 rows (%d)" % len(bills))
check(all(date(2026, 7, 1) <= date.fromisoformat(r["date"]) <= date(2026, 9, 30) for r in inv), "every invoice is dated Jul 1 - Sep 30 2026")
check(all(date(2026, 7, 1) <= date.fromisoformat(r["date"]) <= date(2026, 9, 30) for r in jobs), "every job is dated Jul 1 - Sep 30 2026")

# jobs <-> invoices: every invoice_no referenced by a job exists and sums match, except the planted re-clean
inv_by_no = {r["invoice_no"]: r for r in inv}
job_sum = defaultdict(int)
for j in jobs:
    if j["invoice_no"]:
        job_sum[j["invoice_no"]] += cents(j["price"])
mismatch = [no for no, s in job_sum.items() if no not in inv_by_no or cents(inv_by_no[no]["amount"]) != s]
check(all(no in inv_by_no for no in job_sum), "every invoice_no in jobs-export exists in QuickBooks")
check(mismatch == [m for m in mismatch if inv_by_no[m]["customer"] == "Priya Natarajan" and inv_by_no[m]["amount"] == "185.00"],
      "every invoice equals the sum of its jobs except the planted Priya re-clean (%s)" % ", ".join(mismatch))
# every invoice has at least one job behind it
check(all(no in job_sum for no in inv_by_no), "every QuickBooks invoice has jobs behind it in jobs-export")

# 1. Eastern Promenade credit promised, not applied
tr = open(os.path.join(D02, "meeting-transcript-2026-09-18-ops.txt")).read()
check("$1,200" in tr and "Eastern Promenade" in tr and "three windows" in tr, "transcript promises Eastern Promenade a $1,200 credit for three missed windows")
ep_inv = [r for r in inv if r["customer"] == "Eastern Promenade Condominium Association"]
check(len(ep_inv) >= 3 and all(float(r["amount"]) > 0 and "credit" not in r["description"].lower() for r in ep_inv),
      "no credit line or negative amount on any Eastern Promenade invoice")
ep_sep = [r for r in ep_inv if r["date"].startswith("2026-09") and r["service_line"] == "Commercial cleaning"]
check(len(ep_sep) == 1, "Eastern Promenade has exactly one September cleaning invoice (%s)" % (ep_sep[0]["invoice_no"] if ep_sep else "none"))
wc = len(tr.split())
check(2200 <= wc <= 2900, "transcript is ~2,500 words (%d)" % wc)

# 2. Old Port Dental rate cut not applied
op = {r["date"][:7]: r for r in inv if r["customer"] == "Old Port Dental"}
check(sorted(op) == ["2026-07", "2026-08", "2026-09"] and all(op[m]["amount"] == "640.00" for m in op), "Old Port Dental invoiced $640.00 in Jul, Aug and Sep")
op_txt = pdftotext(os.path.join(D02, "gmail-export", "Re_ Q3 rate change - Old Port Dental.pdf"))
check("$560" in op_txt and "August 1" in op_txt and "$640" in op_txt, "Old Port Dental email: $640 -> $560 effective August 1")

# 3. Munjoy Hill Senior Residence: closed won, never started
wb = load_workbook(os.path.join(D02, "hubspot-deals.xlsx"))
ws = wb["Deals"]
deals = [dict(zip([c.value for c in ws[1]], [c.value for c in row])) for row in ws.iter_rows(min_row=2)]
mh = [d for d in deals if d["Company"] == "Munjoy Hill Senior Residence"]
check(len(mh) == 1 and mh[0]["Deal Stage"] == "Closed Won" and mh[0]["Amount (monthly)"] == 2400 and mh[0]["Close Date"].date() == date(2026, 7, 8),
      "HubSpot: Munjoy Hill Senior Residence Closed Won 2026-07-08 at $2,400/mo")
check(len(deals) == 15, "hubspot-deals has 15 deals")
check(not any("Munjoy Hill Senior" in r["customer"] for r in jobs + inv), "Munjoy Hill Senior Residence has no jobs and no invoices (3 x $2,400 = $7,200)")
# other closed-won deals in Q3 DO have invoices
for name in ("Riverton Family Chiropractic", "Pleasant Hill Veterinary"):
    check(any(r["customer"] == name for r in inv), "closed-won deal %s has invoices (so it reconciles)" % name)

# 4. Back Cove Bakery handyman, completed, never invoiced
bc = [j for j in jobs if j["customer"] == "Back Cove Bakery" and j["service_line"] == "Handyman"]
check(len(bc) == 3 and all(j["status"] == "Completed" and j["invoice_no"] == "" and j["date"].startswith("2026-08") for j in bc)
      and sum(cents(j["price"]) for j in bc) == 197500, "Back Cove Bakery: 3 completed August handyman jobs, blank invoice_no, $1,975.00")
check(not any(r["customer"] == "Back Cove Bakery" and r["service_line"] == "Handyman" for r in inv), "no handyman invoice for Back Cove Bakery in Q3")

# 5. CB-2211 paid twice
cb = [b for b in bills if b["invoice_number"] == "CB-2211"]
check(len(cb) == 2 and all(b["amount"] == "1184.60" for b in cb) and cb[0]["paid_date"] != cb[1]["paid_date"] and cb[0]["bill_no"] != cb[1]["bill_no"],
      "CB-2211 paid twice, $1,184.60 each, on different dates (%s)" % ", ".join(b["paid_date"] for b in cb))
other_dupes = [k for k, n in defaultdict(int, {}).items()]
counts = defaultdict(int)
for b in bills:
    counts[b["invoice_number"]] += 1
check([k for k, n in counts.items() if n > 1] == ["CB-2211"], "CB-2211 is the only vendor invoice paid twice")

# 6. Priya Natarajan free re-clean invoiced
slack = open(os.path.join(D02, "slack-export", "ops-channel.txt")).read()
check(re.search(r"Priya Natarajan.*no charge", slack) is not None, "Slack: crew lead promises Priya Natarajan a re-clean at no charge")
pr = [r for r in inv if r["customer"] == "Priya Natarajan" and r["amount"] == "185.00"]
check(len(pr) == 1 and "re-clean" in pr[0]["description"].lower(), "QuickBooks: Priya Natarajan re-clean invoiced $185.00 (%s)" % (pr[0]["invoice_no"] if pr else "none"))
sl = slack.strip().split("\n")
check(70 <= len(sl) <= 95 and all(re.match(r"\[2026-0[789]-\d\d \d\d:\d\d\] \w+: ", l) for l in sl), "Slack export is ~80 lines in export style (%d)" % len(sl))
check(len([f for f in os.listdir(os.path.join(D02, "gmail-export")) if f.endswith(".pdf")]) == 8, "8 Gmail PDFs")

# ------------------------------------------------------------ 03 dashboard
print("\n[03-dashboard]")
wb = load_workbook(os.path.join(D03, "northwind-revenue-30mo.xlsx"))
check(wb.sheetnames == ["Monthly", "Customers", "Notes"], "sheets are Monthly, Customers, Notes")
ws = wb["Monthly"]
hdr = [c.value for c in ws[1]]
check(hdr[:2] == ["month_end", "month"], "month_end (real date) is column A, month (text) is column B")
xrows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
check(all(hasattr(r[0], "year") for r in xrows) and all(isinstance(r[1], str) and re.match(r"^\d{4}-\d{2}$", r[1]) for r in xrows),
      "month_end cells are real dates, month cells are text yyyy-mm")
csv_rows = read_csv(os.path.join(D03, "revenue_30mo.csv"))
check(len(csv_rows) == len(xrows) and all(
    c["month_end"] == r[0].date().isoformat() and c["customer"] == r[3] and c["service_line"] == r[4] and cents(c["revenue"]) == cents(r[5])
    for c, r in zip(csv_rows, xrows)), "revenue_30mo.csv matches the Monthly sheet row for row (%d rows)" % len(xrows))
months = sorted(set(r[1] for r in xrows))
check(len(months) == 30 and months[0] == "2024-04" and months[-1] == "2026-09", "30 months, 2024-04 to 2026-09")

# THE tie-out: last three months == Q3 QuickBooks invoices, by customer and month (and by service line)
dash = defaultdict(int)
dash_line = defaultdict(int)
for r in xrows:
    if r[1] in ("2026-07", "2026-08", "2026-09"):
        dash[(r[3], r[1])] += cents(r[5])
        dash_line[(r[3], r[1], r[4])] += cents(r[5])
qb = defaultdict(int)
qb_line = defaultdict(int)
for r in inv:
    qb[(r["customer"], r["date"][:7])] += cents(r["amount"])
    qb_line[(r["customer"], r["date"][:7], r["service_line"])] += cents(r["amount"])
check({k: v for k, v in dash.items() if v} == dict(qb), "TIE-OUT: dashboard Jul/Aug/Sep 2026 revenue == QuickBooks invoices, by customer and month, to the cent")
check({k: v for k, v in dash_line.items() if v} == dict(qb_line), "TIE-OUT: same by customer, month and service line")
q3_dash = sum(v for k, v in dash.items())
check(q3_dash == sum(cents(r["amount"]) for r in inv), "TIE-OUT: Q3 total $%s on both sides" % f"{q3_dash/100:,.2f}")

# concentration
ttm = [m for m in months if "2025-10" <= m <= "2026-09"]
ttm_total = sum(cents(r[5]) for r in xrows if r[1] in ttm)
ep_ttm = sum(cents(r[5]) for r in xrows if r[1] in ttm and r[3] == "Eastern Promenade Condominium Association")
share = ep_ttm / ttm_total
check(abs(share - 0.31) < 0.0015, "PLANTED: Eastern Promenade is 31%% of trailing-12-month revenue (%.2f%%)" % (100 * share))
check(2_600_000_00 <= ttm_total <= 3_000_000_00, "trailing-12-month revenue is about $2.8M ($%s)" % f"{ttm_total/100:,.0f}")
others = [k for k in set(r[3] for r in xrows) if k != "Eastern Promenade Condominium Association"]
check(all(sum(cents(r[5]) for r in xrows if r[1] in ttm and r[3] == c) / ttm_total < 0.25 for c in others), "no other customer is over 25% (so only one flag)")

# handyman margin walk
def hand(ms):
    rs = [r for r in xrows if r[4] == "Handyman" and r[1] in ms]
    return sum(r[7] for r in rs) / sum(r[8] for r in rs), sum(r[5] for r in rs) / sum(r[8] for r in rs), 1 - sum(r[6] for r in rs) / sum(r[5] for r in rs)
h0, p0, m0 = hand(months[:3])
h1, p1, m1 = hand(months[-3:])
check(2.0 <= h0 <= 2.3 and 3.25 <= h1 <= 3.55, "PLANTED: handyman crew hours per job rise from ~2.1 (%.2f) to ~3.4 (%.2f)" % (h0, h1))
check(abs(p1 - p0) / p0 < 0.12, "PLANTED: handyman price per job stays roughly flat ($%.0f -> $%.0f)" % (p0, p1))
check(m0 - m1 > 0.12, "PLANTED: handyman gross margin walks down (%.0f%% -> %.0f%%)" % (100 * m0, 100 * m1))

# seasonality
turn = defaultdict(int)
for r in xrows:
    if r[3] in ("Casco Coast Property Management", "Longfellow Street Rentals", "Peaks Island Rentals"):
        turn[r[1][5:]] += cents(r[5])
check(turn["07"] > 2.5 * turn["01"] and turn["08"] > 2.5 * turn["02"], "PLANTED: summer turnover revenue is >2.5x winter (Jul $%s vs Jan $%s)" % (f"{turn['07']/100:,.0f}", f"{turn['01']/100:,.0f}"))
check(not any("Munjoy Hill Senior" in r[3] for r in xrows), "Munjoy Hill Senior Residence is absent from the dashboard too")
bc_row = [r for r in xrows if r[3] == "Back Cove Bakery" and r[4] == "Handyman" and r[1] == "2026-08"]
check(len(bc_row) == 1 and bc_row[0][5] == 0 and bc_row[0][8] == 3, "dashboard shows Back Cove Bakery Aug 2026 handyman: 3 jobs, $0 revenue")

# ------------------------------------------------------------ 01 inbox
print("\n[01-invoices-inbox]")
files = sorted(f for f in os.listdir(D01) if f not in ("README.md", "rules.md"))
check(28 <= len(files) <= 34, "about 30 files in the inbox (%d)" % len(files))
for (vendor, number, d, total, po, fn, lines) in INBOX_INVOICES:
    txt = pdftotext(os.path.join(D01, fn))
    ok = vendor in txt and number in txt and d.strftime("%B %d, %Y") in txt and "{:,.2f}".format(total) in txt and "Northwind Home Services" in txt
    check(ok, "%-26s extracts vendor, number %s, date, total" % (fn, number))
    if po:
        check(re.search(r"PO Number:\s+" + re.escape(po), txt) is not None, "%-26s has PO %s" % (fn, po))
    else:
        check(number == "TP-0893" and "PO Number" not in txt and "3,250.00" in txt, "PLANTED: %s (%s) has NO PO number, $3,250.00" % (fn, number))
a, b = [open(os.path.join(D01, f), "rb").read() for f in DUPLICATE_FILE]
check(a == b and "CB-2211" in pdftotext(os.path.join(D01, DUPLICATE_FILE[1])), "PLANTED: duplicate %s == %s byte for byte, both CB-2211 $1,184.60" % DUPLICATE_FILE)
check(pdftotext(os.path.join(D01, BLANK_SCAN)).strip() == "", "PLANTED: %s extracts no text" % BLANK_SCAN)
note = open(os.path.join(D01, "note to self.txt")).read()
check(not re.search(r"\b(20\d\d|\d{1,2}/\d{1,2}|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\b", note, re.I) and len(note.split()) > 30,
      "PLANTED: note to self.txt has real content and no date")
for png in ("IMG_4502.png", "receipt.png"):
    with open(os.path.join(D01, png), "rb") as f:
        check(f.read(8) == b"\x89PNG\r\n\x1a\n", "%s is a real PNG" % png)
vendors_seen = set(v for (v, *_rest) in INBOX_INVOICES)
check(len(vendors_seen) == 9 and vendors_seen == set(VENDORS), "9 invented vendors")

# ------------------------------------------------------------ 04 prospect
print("\n[04-prospect]")
pfiles = sorted(os.listdir(D04))
pages = [f for f in pfiles if f.endswith(".pdf") and not f.startswith("Re_")]
check(len(pages) == 6, "6 web-page PDFs")
notes = open(os.path.join(D04, "call-notes-first-call.txt")).read()
check("$110K" in notes and "$6,200/mo" in notes and "end of October" in notes and "three vendors" in notes, "PLANTED: call notes: ~$110K/yr budget, decision end of October, three vendors, floor $6,200/mo")
fu = pdftotext(os.path.join(D04, "Re_ follow up - Lakeshore.pdf"))
check("$80,000" in fu and "$110,000" in fu, "PLANTED: follow-up email walks the budget back to closer to $80,000")
jp = pdftotext(os.path.join(D04, "Job posting - Facilities Coordinator - Lakeshore Property Partners.pdf"))
check("scorecard" in jp and "report" in jp, "job posting reveals vendor scorecard and owner reporting")
news = pdftotext(os.path.join(D04, [f for f in pages if f.startswith("News")][0]))
check("$4,500" in news and "common" in news, "news item: $4,500 code-violation fine over common areas")
pc = pdftotext(os.path.join(D04, "Pricing - ProClean Maine.pdf"))
check("$38" in pc and "per unit" in pc, "competitor pricing page has per-unit monthly pricing")
for f in pages:
    t = pdftotext(os.path.join(D04, f))
    check(".example" in t and "Saved Sep" in t, "%s has a fake URL header and saved date" % f[:50])

# ------------------------------------------------------------ context files
print("\n[context files]")
CONTEXT_FILES = [os.path.join(D01, "rules.md"), os.path.join(D02, "rules.md"),
                 os.path.join(D03, "brand-guidelines.md"), os.path.join(D04, "northwind-company-profile.md")]
ctx = {}
for p in CONTEXT_FILES:
    ok = os.path.isfile(p) and os.path.getsize(p) > 200
    check(ok, "%s exists and is non-empty" % os.path.relpath(p, ROOT))
    ctx[p] = open(p).read() if ok else ""
r01, r02, brand, profile = [ctx[p] for p in CONTEXT_FILES]

# every vendor actually printed in an invoice PDF appears in the rules.md category map (a table row)
map_rows = [l for l in r01.splitlines() if l.startswith("| ") and "`" in l]
pdf_vendors = set()
for f in sorted(os.listdir(D01)):
    if f.endswith(".pdf"):
        lines = [l.strip() for l in pdftotext(os.path.join(D01, f)).splitlines() if l.strip()]
        if lines:
            pdf_vendors.add(re.split(r"\s{2,}", lines[0])[0])
check(pdf_vendors == set(VENDORS), "vendors read from the invoice PDFs are the 9 known vendors (%d)" % len(pdf_vendors))
missing = [v for v in sorted(pdf_vendors) if not any(("| " + v + " |") in l for l in map_rows)]
check(missing == [], "every invoice vendor is in the 01 rules.md category map (%s)" % (missing or "all 9"))
for need in ("_receipts/", "_review/", "Needs a human", "date,vendor,category,invoice_number,po_number,amount,new_path", "YYYY-MM-DD_vendor-slug_invoice-number.pdf"):
    check(need in r01, "01 rules.md specifies %s" % need)
for need in ("calculator", "July 1 to September 30, 2026", "stop", "Closed Won", "duplicate payments", "formulas"):
    check(need in r02, "02 rules.md covers '%s'" % need)
hexes = set(re.findall(r"#[0-9A-Fa-f]{6}\b", brand))
check(len(hexes) >= 6, "brand-guidelines.md has hex color codes (%d)" % len(hexes))
for need in ("Fraunces", "Source Sans 3", "IBM Plex Mono", "pie or donut"):
    check(need in brand, "brand-guidelines.md names %s" % need)
check("6,200" not in profile and "6200" not in profile, "company profile does NOT contain the $6,200 floor")
for need in ("Dana Whitfield", "Eastern Promenade Condominium Association", "$400 per missed window", "Greenline Lawn & Snow", "60-day pilot"):
    check(need in profile, "company profile states '%s'" % need)
dashes = [os.path.relpath(p, ROOT) for p in CONTEXT_FILES + [os.path.join(ROOT, d, "README.md") for d in ("", "01-invoices-inbox", "02-sources", "03-dashboard", "04-prospect")]
          if os.path.exists(p) and ("\u2014" in open(p).read() or "\u2013" in open(p).read())]
check(dashes == [], "no em or en dashes in the context files or READMEs (%s)" % (dashes or "clean"))
SKILLS = {"01-invoices-inbox": ("invoice-inbox", "rules.md"), "02-sources": ("quarterly-reconciliation", "rules.md"),
          "03-dashboard": ("monthly-dashboard", "brand-guidelines.md"), "04-prospect": ("call-prep", "northwind-company-profile.md")}
for d, (name, cf) in SKILLS.items():
    rd = open(os.path.join(ROOT, d, "README.md")).read()
    check(("into a skill called %s," % name) in rd and ("Put everything from %s inside the skill" % cf) in rd and ("Reference " + cf) in rd.replace("reference " + cf, "Reference " + cf),
          "%s/README.md has the skills prompt (%s, %s) and a prompt that references %s" % (d, name, cf, cf))

# ------------------------------------------------------------ hygiene
print("\n[hygiene]")
# stored reversed so a plain grep of the repo for these terms finds nothing, including this file
FORBIDDEN = [t[::-1] for t in ["tniopeslup", "opy", "ramhcrok", "rehtael robrah", "redlac", "nevah", "htraeh", "xofmialc", "sdpi"]]
hits = []
scanned = set()
for dirpath, dirnames, filenames in os.walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in (".git", "__pycache__")]
    for fn in filenames:
        p = os.path.join(dirpath, fn)
        scanned.add(p)
        blobs = [open(p, "rb").read().lower()]
        if fn.lower().endswith(".pdf"):
            blobs.append(pdftotext(p).lower().encode())
        if fn.lower().endswith(".xlsx"):
            with zipfile.ZipFile(p) as z:
                blobs += [z.read(n).lower() for n in z.namelist()]
        for term in FORBIDDEN:
            if any(term.encode() in b for b in blobs):
                hits.append((os.path.relpath(p, ROOT), term))
check(hits == [], "no forbidden terms anywhere (%s)" % (hits if hits else "clean"))
check(all(p in scanned for p in CONTEXT_FILES), "forbidden-term scan covered all 4 context files")

for xl in (os.path.join(D02, "hubspot-deals.xlsx"), os.path.join(D03, "northwind-revenue-30mo.xlsx")):
    with zipfile.ZipFile(xl) as z:
        names = z.namelist()
        bad = [n for n in names if "comments" in n.lower() or "threadedcomment" in n.lower() or "vmldrawing" in n.lower()]
        core = z.read("docProps/core.xml").decode()
        creators = re.findall(r"<dc:creator>(.*?)</dc:creator>|<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core)
        names_in = set(a or b for a, b in creators)
        check(bad == [] and names_in == {"Northwind"}, "%s: no comments/threaded comments/VML; author is Northwind only" % os.path.basename(xl))

total = 0
for dirpath, dirnames, filenames in os.walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in (".git", "__pycache__")]
    total += sum(os.path.getsize(os.path.join(dirpath, f)) for f in filenames)
check(total < 1.5 * 1024 * 1024, "total size under 1.5 MB (%.2f MB)" % (total / 1024 / 1024))

# READMEs cite the real IDs
r02 = open(os.path.join(D02, "README.md")).read() if os.path.exists(os.path.join(D02, "README.md")) else ""
check(pr and pr[0]["invoice_no"] in r02 and ep_sep and ep_sep[0]["invoice_no"] in r02 and all(j["job_id"] in r02 for j in bc) and all(b["bill_no"] in r02 for b in cb),
      "02-sources/README.md cites the real invoice, job and bill numbers")
check(all(os.path.exists(os.path.join(ROOT, d, "README.md")) for d in ("", "01-invoices-inbox", "02-sources", "03-dashboard", "04-prospect")), "every folder has a README")

print("\n%d checks passed, %d failed" % (PASSES, len(FAILS)))
for f in FAILS:
    print("  FAIL " + f)
sys.exit(1 if FAILS else 0)
